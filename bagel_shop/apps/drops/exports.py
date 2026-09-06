from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="22282D")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _excel_safe(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def workbook_response(filename, sheets):
    """Return a styled .xlsx response from (title, headers, rows) sheet tuples."""
    workbook = Workbook()
    workbook.remove(workbook.active)

    for title, headers, rows in sheets:
        worksheet = workbook.create_sheet(title=title[:31])
        worksheet.append(headers)
        for cell in worksheet[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(vertical="center")
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

        for row in rows:
            worksheet.append([_excel_safe(value) for value in row])

        for column_index, header in enumerate(headers, start=1):
            longest = len(str(header))
            if worksheet.max_row >= 2:
                for column in worksheet.iter_cols(
                    min_col=column_index,
                    max_col=column_index,
                    min_row=2,
                    max_row=worksheet.max_row,
                ):
                    for item in column:
                        longest = max(longest, len(str(item.value or "")))
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(
                max(longest + 2, 12), 42
            )

    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
