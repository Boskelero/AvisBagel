from datetime import timedelta
from io import BytesIO

from django.core.management import call_command
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from django.contrib.auth import get_user_model
from django.core import mail
from openpyxl import load_workbook

from bagel_shop.apps.catalog.models import Category, Product
from bagel_shop.apps.blog.models import Post
from bagel_shop.apps.orders.models import Order
from bagel_shop.apps.notifications.models import DropAlert, DropAnnouncement, NewsletterSubscriber
from bagel_shop.apps.notifications.services import notify_owner_about_drop_capacity

from .models import BagelPriceTier, Drop, DropProduct, DropReservation
from .services import calculate_bagel_base_price


@override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
class DropOrderingTests(TestCase):
    def setUp(self):
        BagelPriceTier.objects.all().delete()
        BagelPriceTier.objects.bulk_create([
            BagelPriceTier(quantity=1, price_cents=1200),
            BagelPriceTier(quantity=6, price_cents=6500),
            BagelPriceTier(quantity=12, price_cents=12000),
        ])
        category = Category.objects.create(name="Bagels", name_en="Bagels", slug="bagels-test")
        self.plain = Product.objects.create(
            category=category,
            name="Plain",
            name_en="Plain",
            slug="plain-test",
            price_cents=1200,
            product_type=Product.TYPE_BAGEL,
        )
        self.jalapeno = Product.objects.create(
            category=category,
            name="Jalapeño Cheddar",
            name_en="Jalapeño Cheddar",
            slug="jalapeno-test",
            price_cents=1500,
            specialty_upcharge_cents=300,
            product_type=Product.TYPE_BAGEL,
        )
        now = timezone.now()
        self.drop = Drop.objects.create(
            name="Thursday Drop",
            status=Drop.STATUS_PUBLISHED,
            opens_at=now - timedelta(hours=1),
            closes_at=now + timedelta(hours=2),
            pickup_starts_at=now + timedelta(hours=3),
            pickup_ends_at=now + timedelta(hours=5),
            pickup_location="Ben Yehuda 69, Tel Aviv",
            bagel_capacity=72,
            max_bagels_per_order=12,
        )
        DropProduct.objects.create(drop=self.drop, product=self.plain)
        DropProduct.objects.create(drop=self.drop, product=self.jalapeno)

    def _add(self, client, product, quantity):
        return client.post(reverse("cart:add_product", args=[product.id]), {"quantity": quantity})

    def _checkout(self, client):
        return client.post(
            reverse("checkout:checkout"),
            {
                "customer_name": "Test Customer",
                "email": "test@example.com",
                "phone": "0501234567",
                "payment_method": Order.PAYMENT_PAY_ON_PICKUP,
                "notes": "",
            },
        )

    def test_tiered_pricing_uses_best_combination(self):
        self.assertEqual(calculate_bagel_base_price(5), 6000)
        self.assertEqual(calculate_bagel_base_price(6), 6500)
        self.assertEqual(calculate_bagel_base_price(8), 8900)
        self.assertEqual(calculate_bagel_base_price(12), 12000)
        self.assertEqual(calculate_bagel_base_price(18), 18500)

    def test_six_pack_adds_only_specialty_upcharge(self):
        client = Client()
        self._add(client, self.plain, 5)
        self._add(client, self.jalapeno, 1)
        response = self._checkout(client)
        self.assertEqual(response.status_code, 302)
        order = Order.objects.get()
        self.assertEqual(order.bagel_quantity, 6)
        self.assertEqual(order.total_cents, 6800)
        self.assertEqual(order.discount_cents, 700)
        self.assertEqual(order.capacity_reservation.status, DropReservation.STATUS_CONFIRMED)

    def test_order_confirmation_is_private_to_customer_session(self):
        client = Client()
        self._add(client, self.plain, 1)
        self._checkout(client)
        order = Order.objects.get()
        url = reverse("checkout:success", args=[order.number])
        self.assertEqual(client.get(url).status_code, 200)
        self.assertEqual(Client().get(url).status_code, 404)

    def test_order_cap_is_enforced_at_checkout(self):
        client = Client()
        self._add(client, self.plain, 13)
        response = self._checkout(client)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "maximum of 12 bagels")
        self.assertFalse(Order.objects.exists())

    def test_shared_capacity_cannot_be_oversold(self):
        self.drop.bagel_capacity = 6
        self.drop.max_bagels_per_order = 6
        self.drop.save()
        first = Client()
        second = Client()
        self._add(first, self.plain, 6)
        self._add(second, self.plain, 1)
        self.assertEqual(self._checkout(first).status_code, 302)

        response = self._checkout(second)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "no longer accepting orders")
        self.assertEqual(Order.objects.count(), 1)
        self.assertEqual(self.drop.remaining_bagels, 0)
        home = Client().get(reverse("pages:home"))
        self.assertContains(home, "This Drop sold out")
        self.assertNotContains(home, "Ordering opens in")

    def test_canceling_order_releases_capacity(self):
        client = Client()
        self._add(client, self.plain, 6)
        self._checkout(client)
        order = Order.objects.get()
        order.status = Order.STATUS_CANCELED
        order.save()
        order.capacity_reservation.refresh_from_db()
        self.assertEqual(order.capacity_reservation.status, DropReservation.STATUS_RELEASED)
        self.assertEqual(self.drop.remaining_bagels, 72)

    def test_product_outside_drop_cannot_be_added(self):
        extra = Product.objects.create(
            category=self.plain.category,
            name="Future flavor",
            slug="future-test",
            price_cents=1200,
        )
        response = self._add(Client(), extra, 1)
        self.assertEqual(response.status_code, 302)
        self.assertFalse(Order.objects.exists())

    def test_products_outside_the_drop_are_presented_as_rotating_menu_items(self):
        future = Product.objects.create(
            category=self.plain.category,
            name="Future flavor",
            name_en="Future flavor",
            slug="future-menu-test",
            price_cents=1200,
        )
        menu = Client().get(reverse("catalog:menu"))
        self.assertContains(menu, "Available in selected Drops")
        self.assertNotContains(menu, ">Unavailable<")

        detail = Client().get(reverse("catalog:product_detail", args=[future.slug]))
        self.assertContains(detail, "Part of our rotating menu")
        self.assertContains(detail, "Get Drop alerts")
        self.assertNotContains(detail, "not available in the current Drop")

    def test_drop_home_and_menu_render(self):
        response = Client().get(reverse("pages:home"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Thursday Drop")
        self.assertContains(response, "The Drop is live")
        self.assertNotContains(response, "bagel-header-1.jpg")
        response = Client().get(reverse("catalog:menu"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Ordering is open")

    def test_mobile_order_builder_creates_drop_cart(self):
        client = Client()
        page = client.get(reverse("drops:order"))
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "72 bagels remain")
        response = client.post(
            reverse("drops:order"),
            {
                f"product_{self.plain.id}": 5,
                f"product_{self.jalapeno.id}": 1,
            },
        )
        self.assertRedirects(response, reverse("cart:detail"))
        cart = client.session["cart"]
        self.assertEqual(cart["drop_id"], self.drop.id)
        self.assertEqual(sum(line["quantity"] for line in cart["lines"]), 6)

    def test_staff_dashboard_requires_staff_and_can_duplicate_drop(self):
        anonymous = Client()
        response = anonymous.get(reverse("drops:staff_dashboard"))
        self.assertEqual(response.status_code, 302)

        staff = get_user_model().objects.create_user(
            username="avi", password="test-password", is_staff=True
        )
        client = Client()
        client.force_login(staff)
        self.assertEqual(client.get(reverse("drops:staff_dashboard")).status_code, 200)
        self.assertEqual(client.get(reverse("drops:staff_orders")).status_code, 200)
        self.assertEqual(client.get(reverse("drops:staff_customers")).status_code, 200)
        self.assertEqual(client.get(reverse("drops:staff_products")).status_code, 200)
        self.assertEqual(client.get(reverse("drops:staff_blog_posts")).status_code, 200)
        self.assertEqual(client.get(reverse("drops:staff_categories")).status_code, 200)
        self.assertEqual(client.get(reverse("drops:staff_pricing")).status_code, 200)
        self.assertEqual(client.get(reverse("drops:staff_seo")).status_code, 200)
        self.assertEqual(client.get(reverse("drops:staff_inquiries")).status_code, 200)
        self.assertEqual(client.get(reverse("drops:staff_subscribers")).status_code, 200)
        for export_name in [
            "staff_orders_excel",
            "staff_customers_excel",
            "staff_products_excel",
            "staff_inquiries_excel",
            "staff_subscribers_excel",
        ]:
            export = client.get(reverse(f"drops:{export_name}"))
            self.assertEqual(export.status_code, 200)
            self.assertTrue(export.content.startswith(b"PK"))
        self.assertEqual(client.get(reverse("drops:staff_edit", args=[self.drop.id])).status_code, 200)
        self.assertEqual(client.get(reverse("drops:staff_production_sheet", args=[self.drop.id])).status_code, 200)
        self.assertTrue(
            client.get(reverse("drops:staff_excel", args=[self.drop.id])).content.startswith(b"PK")
        )
        response = client.post(reverse("drops:staff_duplicate", args=[self.drop.id]))
        duplicate = Drop.objects.exclude(pk=self.drop.pk).get()
        self.assertRedirects(response, reverse("drops:staff_edit", args=[duplicate.id]))
        self.assertEqual(duplicate.status, Drop.STATUS_DRAFT)
        self.assertEqual(duplicate.product_entries.count(), 2)
        self.assertEqual(duplicate.opens_at, self.drop.opens_at + timedelta(days=7))

    def test_staff_can_delete_an_empty_draft_drop(self):
        draft = Drop.objects.create(
            name="Accidental Drop",
            status=Drop.STATUS_DRAFT,
            opens_at=self.drop.opens_at + timedelta(days=7),
            closes_at=self.drop.closes_at + timedelta(days=7),
            pickup_starts_at=self.drop.pickup_starts_at + timedelta(days=7),
            pickup_ends_at=self.drop.pickup_ends_at + timedelta(days=7),
            pickup_location=self.drop.pickup_location,
            bagel_capacity=75,
            max_bagels_per_order=4,
        )
        DropProduct.objects.create(drop=draft, product=self.plain)
        staff = get_user_model().objects.create_user(username="drop-deleter", is_staff=True)
        client = Client()
        client.force_login(staff)

        confirmation = client.get(reverse("drops:staff_delete", args=[draft.id]))
        self.assertContains(confirmation, "permanently delete")
        response = client.post(reverse("drops:staff_delete", args=[draft.id]))

        self.assertRedirects(response, reverse("drops:staff_dashboard"))
        self.assertFalse(Drop.objects.filter(pk=draft.id).exists())

    def test_staff_lists_are_paginated_and_excel_exports_open(self):
        NewsletterSubscriber.objects.bulk_create([
            NewsletterSubscriber(email=f"person{index:02d}@example.com")
            for index in range(30)
        ])
        staff = get_user_model().objects.create_user(username="exporter", is_staff=True)
        client = Client()
        client.force_login(staff)

        listing = client.get(reverse("drops:staff_subscribers"), {"page": 2})
        self.assertEqual(listing.context["page_obj"].paginator.num_pages, 2)
        self.assertEqual(len(listing.context["subscribers"]), 5)

        export = client.get(reverse("drops:staff_subscribers_excel"))
        self.assertEqual(
            export["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(export.content), read_only=True)
        self.assertEqual(workbook["Subscribers"].max_row, 31)

    def test_flash_messages_are_closable_and_auto_dismiss(self):
        staff = get_user_model().objects.create_user(username="messenger", is_staff=True)
        client = Client()
        client.force_login(staff)
        response = client.post(
            reverse("drops:staff_action", args=[self.drop.id]),
            {"action": "extend", "hours": 1},
            follow=True,
        )
        self.assertContains(response, 'data-auto-dismiss="6000"')
        self.assertContains(response, "btn-close")

    def test_demo_seed_is_complete_and_idempotent(self):
        call_command("seed_demo_data", verbosity=0)
        call_command("seed_demo_data", verbosity=0)
        self.assertEqual(Product.objects.filter(slug__in=[
            "plain", "sesame", "poppy", "onion", "everything", "mediterranean",
            "zaatar", "cinnamon-raisin", "jalapeno-cheddar",
        ]).count(), 9)
        self.assertTrue(Category.objects.filter(slug="cream-cheese", is_active=True).exists())
        self.assertTrue(Category.objects.filter(slug="jams", is_active=True).exists())
        self.assertEqual(
            Product.objects.filter(
                slug__in=[
                    "plain-cream-cheese", "scallion-cream-cheese",
                    "strawberry-jam", "seasonal-jam",
                ],
                product_type=Product.TYPE_EXTRA,
            ).count(),
            4,
        )
        self.assertEqual(Order.objects.filter(number__startswith="DEMO").count(), 20)
        self.assertEqual(
            NewsletterSubscriber.objects.filter(email__startswith="demo.customer").count(), 20
        )
        self.assertEqual(Drop.objects.filter(name="Demo Bagel Drop").count(), 1)
        self.assertEqual(Post.objects.filter(slug__startswith="demo-").count(), 3)

    def test_staff_cannot_delete_a_published_drop(self):
        staff = get_user_model().objects.create_user(username="careful-deleter", is_staff=True)
        client = Client()
        client.force_login(staff)

        confirmation = client.get(reverse("drops:staff_delete", args=[self.drop.id]))
        self.assertContains(confirmation, "Close this Drop before deleting it")
        response = client.post(
            reverse("drops:staff_delete", args=[self.drop.id]),
            follow=True,
        )

        self.assertTrue(Drop.objects.filter(pk=self.drop.id).exists())
        self.assertContains(response, "Close this Drop before deleting it")

    def test_staff_can_add_product_using_shekel_prices(self):
        staff = get_user_model().objects.create_user(username="product-owner", is_staff=True)
        client = Client()
        client.force_login(staff)
        response = client.post(
            reverse("drops:staff_product_create"),
            {
                "category": self.plain.category_id,
                "name_en": "Weekly Special",
                "name_he": "",
                "description_en": "A rotating special.",
                "description_he": "",
                "product_type": Product.TYPE_BAGEL,
                "price_ils": "14.50",
                "specialty_upcharge_ils": "2.50",
                "is_active": "on",
            },
        )
        self.assertRedirects(response, reverse("drops:staff_products"))
        product = Product.objects.get(name_en="Weekly Special")
        self.assertEqual(product.price_cents, 1450)
        self.assertEqual(product.specialty_upcharge_cents, 250)

    def test_staff_can_create_and_publish_a_blog_post(self):
        staff = get_user_model().objects.create_user(username="blog-owner", is_staff=True)
        client = Client()
        client.force_login(staff)
        published_at = timezone.localtime(timezone.now()).replace(second=0, microsecond=0)
        response = client.post(
            reverse("drops:staff_blog_post_create"),
            {
                "category": "",
                "title_en": "A fresh test post",
                "excerpt_en": "News from the bakery.",
                "content_en": "This is the full story from the bakery.",
                "title_he": "",
                "excerpt_he": "",
                "content_he": "",
                "status": Post.STATUS_PUBLISHED,
                "published_at": published_at.strftime("%Y-%m-%dT%H:%M"),
            },
        )
        post = Post.objects.get(title_en="A fresh test post")
        self.assertRedirects(response, reverse("drops:staff_blog_posts"))
        self.assertEqual(
            client.get(reverse("drops:staff_blog_post_edit", args=[post.id])).status_code,
            200,
        )
        self.assertEqual(Client().get(post.get_absolute_url()).status_code, 200)

    def test_staff_can_update_bagel_deal_prices(self):
        staff = get_user_model().objects.create_user(username="pricing-owner", is_staff=True)
        client = Client()
        client.force_login(staff)
        tiers = list(BagelPriceTier.objects.order_by("quantity"))
        data = {
            "form-TOTAL_FORMS": len(tiers),
            "form-INITIAL_FORMS": len(tiers),
            "form-MIN_NUM_FORMS": 0,
            "form-MAX_NUM_FORMS": 1000,
        }
        for index, tier in enumerate(tiers):
            data.update({
                f"form-{index}-id": tier.id,
                f"form-{index}-quantity": tier.quantity,
                f"form-{index}-price_ils": "64.00" if tier.quantity == 6 else f"{tier.price_cents / 100:.2f}",
                f"form-{index}-is_active": "on",
            })
        response = client.post(reverse("drops:staff_pricing"), data)
        self.assertRedirects(response, reverse("drops:staff_pricing"))
        self.assertEqual(BagelPriceTier.objects.get(quantity=6).price_cents, 6400)

    def test_staff_can_send_one_drop_announcement(self):
        NewsletterSubscriber.objects.create(email="subscriber@example.com")
        staff = get_user_model().objects.create_user(username="announcer", is_staff=True)
        client = Client()
        client.force_login(staff)
        url = reverse("drops:staff_announcement", args=[self.drop.id])
        self.assertEqual(client.get(url).status_code, 200)
        self.assertRedirects(client.post(url), reverse("drops:staff_detail", args=[self.drop.id]))
        self.assertEqual(DropAnnouncement.objects.get(drop=self.drop).recipient_count, 1)
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn("/en/drops/order/", mail.outbox[0].body)
        client.post(url)
        self.assertEqual(DropAnnouncement.objects.filter(drop=self.drop).count(), 1)

    def test_scheduled_drop_announcement_links_to_homepage_countdown(self):
        self.drop.opens_at = timezone.now() + timedelta(days=1)
        self.drop.closes_at = timezone.now() + timedelta(days=2)
        self.drop.pickup_starts_at = timezone.now() + timedelta(days=2, hours=1)
        self.drop.pickup_ends_at = timezone.now() + timedelta(days=2, hours=3)
        self.drop.save()
        NewsletterSubscriber.objects.create(email="waiting@example.com")
        staff = get_user_model().objects.create_user(username="prelaunch-announcer", is_staff=True)
        client = Client()
        client.force_login(staff)

        homepage = client.get(reverse("pages:home"))
        self.assertContains(homepage, "Ordering starts in")
        self.assertContains(homepage, "Email me before the Drop")
        self.assertNotContains(homepage, "bagel-header-1.jpg")

        client.post(reverse("drops:staff_announcement", args=[self.drop.id]))
        self.assertIn("See the countdown:", mail.outbox[0].body)
        self.assertNotIn("/en/drops/order/", mail.outbox[0].body)

    def test_staff_cancel_action_releases_capacity(self):
        customer = Client()
        self._add(customer, self.plain, 6)
        self._checkout(customer)
        order = Order.objects.get()
        staff = get_user_model().objects.create_user(username="owner", is_staff=True)
        client = Client()
        client.force_login(staff)
        response = client.post(
            reverse("drops:staff_order_status", args=[self.drop.id, order.id]),
            {"status": Order.STATUS_CANCELED},
        )
        self.assertRedirects(response, reverse("drops:staff_detail", args=[self.drop.id]))
        order.capacity_reservation.refresh_from_db()
        self.assertEqual(order.capacity_reservation.status, DropReservation.STATUS_RELEASED)

    def test_staff_order_workspace_replaces_admin_order_link(self):
        customer = Client()
        self._add(customer, self.plain, 2)
        self._checkout(customer)
        order = Order.objects.get()
        staff = get_user_model().objects.create_user(username="order-manager", is_staff=True)
        client = Client()
        client.force_login(staff)

        listing = client.get(reverse("drops:staff_orders"), {"q": order.number})
        self.assertContains(listing, order.number)
        self.assertContains(listing, reverse("drops:staff_order_detail", args=[order.id]))
        detail = client.get(reverse("drops:staff_order_detail", args=[order.id]))
        self.assertContains(detail, "Test Customer")
        self.assertContains(detail, "2 × Plain")

        response = client.post(
            reverse("drops:staff_order_status", args=[self.drop.id, order.id]),
            {"status": Order.STATUS_READY, "return": "order"},
        )
        self.assertRedirects(response, reverse("drops:staff_order_detail", args=[order.id]))
        order.refresh_from_db()
        self.assertEqual(order.status, Order.STATUS_READY)

    def test_checkout_get_redirects_when_cart_exceeds_order_cap(self):
        client = Client()
        self._add(client, self.plain, 13)
        response = client.get(reverse("checkout:checkout"))
        self.assertRedirects(response, reverse("cart:detail"))
        cart = client.get(reverse("cart:detail"))
        self.assertContains(cart, "maximum of 12 bagels")

    def test_canceled_order_cannot_be_reactivated_if_capacity_was_resold(self):
        self.drop.bagel_capacity = 6
        self.drop.max_bagels_per_order = 6
        self.drop.save()
        first_customer = Client()
        self._add(first_customer, self.plain, 6)
        self._checkout(first_customer)
        first_order = Order.objects.get()
        first_order.status = Order.STATUS_CANCELED
        first_order.save()

        second_customer = Client()
        self._add(second_customer, self.plain, 6)
        self._checkout(second_customer)

        staff = get_user_model().objects.create_user(username="capacity-manager", is_staff=True)
        client = Client()
        client.force_login(staff)
        response = client.post(
            reverse("drops:staff_order_status", args=[self.drop.id, first_order.id]),
            {"status": Order.STATUS_READY, "return": "order"},
            follow=True,
        )
        first_order.refresh_from_db()
        self.assertEqual(first_order.status, Order.STATUS_CANCELED)
        self.assertContains(response, "cannot be reactivated")
        self.assertEqual(self.drop.remaining_bagels, 0)

    def test_drop_cannot_be_published_without_an_active_bagel(self):
        self.drop.status = Drop.STATUS_DRAFT
        self.drop.save()
        self.drop.product_entries.all().delete()
        staff = get_user_model().objects.create_user(username="publisher", is_staff=True)
        client = Client()
        client.force_login(staff)
        response = client.post(
            reverse("drops:staff_action", args=[self.drop.id]),
            {"action": "publish"},
            follow=True,
        )
        self.drop.refresh_from_db()
        self.assertEqual(self.drop.status, Drop.STATUS_DRAFT)
        self.assertContains(response, "Add at least one active bagel")

    def test_staff_can_create_drop_with_selected_products(self):
        staff = get_user_model().objects.create_user(username="creator", is_staff=True)
        client = Client()
        client.force_login(staff)
        start = timezone.localtime(timezone.now() + timedelta(days=14)).replace(second=0, microsecond=0)
        response = client.post(
            reverse("drops:staff_create"),
            {
                "name": "Future Drop",
                "status": Drop.STATUS_DRAFT,
                "opens_at": start.strftime("%Y-%m-%dT%H:%M"),
                "closes_at": (start + timedelta(days=1)).strftime("%Y-%m-%dT%H:%M"),
                "pickup_starts_at": (start + timedelta(days=1, hours=1)).strftime("%Y-%m-%dT%H:%M"),
                "pickup_ends_at": (start + timedelta(days=1, hours=3)).strftime("%Y-%m-%dT%H:%M"),
                "pickup_location": "Ben Yehuda 69, Tel Aviv",
                "bagel_capacity": 90,
                "max_bagels_per_order": 12,
                "products": [self.plain.id],
            },
        )
        created = Drop.objects.get(name="Future Drop")
        self.assertRedirects(response, reverse("drops:staff_detail", args=[created.id]))
        self.assertEqual(list(created.products.all()), [self.plain])

    def test_capacity_alerts_are_deduplicated(self):
        self.drop.bagel_capacity = 10
        self.drop.max_bagels_per_order = 10
        self.drop.save()
        client = Client()
        self._add(client, self.plain, 6)
        self._checkout(client)
        notify_owner_about_drop_capacity(self.drop)
        notify_owner_about_drop_capacity(self.drop)
        self.assertEqual(
            DropAlert.objects.filter(alert_type=DropAlert.TYPE_LOW_CAPACITY).count(), 1
        )
